from __future__ import annotations

import shutil
from pathlib import Path

from 主程序.ner引擎 import 全局映射表


def 脱敏并保存原格式(文件路径: str, 映射: 全局映射表, 输出目录: str | None = None) -> str:
    后缀 = Path(文件路径).suffix.lower()
    输出路径 = _构建输出路径(文件路径, 输出目录)
    if 后缀 in {".txt", ".md"}:
        return _脱敏文本文件(文件路径, 映射, 输出路径)
    if 后缀 == ".docx":
        return _脱敏docx(文件路径, 映射, 输出路径)
    if 后缀 == ".xlsx":
        return _脱敏excel(文件路径, 映射, 输出路径)
    if 后缀 == ".pptx":
        return _脱敏ppt(文件路径, 映射, 输出路径)
    if 后缀 == ".pdf":
        return _脱敏pdf(文件路径, 映射, 输出路径)
    raise ValueError(f"不支持的格式：{后缀}")


def 还原并保存原格式(文件路径: str, 映射: 全局映射表, 输出目录: str | None = None) -> str:
    后缀 = Path(文件路径).suffix.lower()
    输出路径 = _构建还原路径(文件路径, 输出目录)
    if 后缀 in {".txt", ".md"}:
        return _还原文本文件(文件路径, 映射, 输出路径)
    if 后缀 == ".docx":
        return _还原docx(文件路径, 映射, 输出路径)
    if 后缀 == ".xlsx":
        return _还原excel(文件路径, 映射, 输出路径)
    if 后缀 == ".pptx":
        return _还原ppt(文件路径, 映射, 输出路径)
    if 后缀 == ".pdf":
        return _还原pdf(文件路径, 映射, 输出路径)
    raise ValueError(f"不支持的格式：{后缀}")


def 生成脱敏md(文件路径: str, 映射: 全局映射表, 输出目录: str | None = None,
             原始文本: str | None = None) -> str:
    # 扫描版PDF没有文字层，调用方应传入OCR出的文本；未传入时才从文件提取
    if 原始文本 is None:
        from 主程序.文档解析器 import 提取文本
        原始文本 = 提取文本(文件路径)
    脱敏文本 = 映射.批量替换文本(原始文本)
    输出路径 = _构建md输出路径(文件路径, 输出目录)
    Path(输出路径).write_text(脱敏文本, encoding="utf-8")
    return 输出路径


def _构建输出路径(文件路径: str, 输出目录: str | None) -> str:
    原始 = Path(文件路径)
    目标 = Path(输出目录) if 输出目录 else 原始.parent
    目标.mkdir(parents=True, exist_ok=True)
    return str(目标 / f"{原始.stem}_脱敏{原始.suffix}")


def _构建还原路径(文件路径: str, 输出目录: str | None) -> str:
    原始 = Path(文件路径)
    目标 = Path(输出目录) if 输出目录 else 原始.parent
    目标.mkdir(parents=True, exist_ok=True)
    return str(目标 / f"{原始.stem}_还原{原始.suffix}")


def _构建md输出路径(文件路径: str, 输出目录: str | None) -> str:
    原始 = Path(文件路径)
    目标 = Path(输出目录) if 输出目录 else 原始.parent
    目标.mkdir(parents=True, exist_ok=True)
    return str(目标 / f"{原始.stem}_脱敏.md")


def _脱敏文本文件(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    文本 = Path(路径).read_text(encoding="utf-8")
    脱敏文本 = 映射.批量替换文本(文本)
    Path(输出路径).write_text(脱敏文本, encoding="utf-8")
    return 输出路径


def _还原文本文件(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    文本 = Path(路径).read_text(encoding="utf-8")
    还原文本 = 映射.批量还原文本(文本)
    Path(输出路径).write_text(还原文本, encoding="utf-8")
    return 输出路径


def _脱敏docx(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    from docx import Document
    from 主程序._docx_xml工具 import 收集文档全部部件, 替换部件文本
    from 主程序.文档解析器 import 验证office文件格式
    验证office文件格式(路径)
    shutil.copy2(路径, 输出路径)
    doc = Document(输出路径)
    for 部件 in 收集文档全部部件(doc):
        替换部件文本(部件, 映射.批量替换文本)
    doc.save(输出路径)
    return 输出路径


def _还原docx(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    from docx import Document
    from 主程序._docx_xml工具 import 收集文档全部部件, 替换部件文本
    from 主程序.文档解析器 import 验证office文件格式
    验证office文件格式(路径)
    shutil.copy2(路径, 输出路径)
    doc = Document(输出路径)
    for 部件 in 收集文档全部部件(doc):
        替换部件文本(部件, 映射.批量还原文本)
    doc.save(输出路径)
    return 输出路径


def _脱敏excel(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    from openpyxl import load_workbook
    from 主程序.文档解析器 import 验证office文件格式
    验证office文件格式(路径)
    shutil.copy2(路径, 输出路径)
    wb = load_workbook(输出路径)
    for 工作表 in wb.worksheets:
        for 行 in 工作表.iter_rows():
            for 单元格 in 行:
                val = 单元格.value
                if val is None:
                    continue
                if isinstance(val, (int, float)):
                    if isinstance(val, float) and val == int(val) and abs(val) >= 1e15:
                        文本 = f"{int(val)}"
                    else:
                        continue
                else:
                    文本 = str(val)
                替换结果 = 映射.批量替换文本(文本)
                if 替换结果 != 文本:
                    单元格.value = 替换结果
    wb.save(输出路径)
    wb.close()
    return 输出路径


def _还原excel(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    from openpyxl import load_workbook
    from 主程序.文档解析器 import 验证office文件格式
    验证office文件格式(路径)
    shutil.copy2(路径, 输出路径)
    wb = load_workbook(输出路径)
    for 工作表 in wb.worksheets:
        for 行 in 工作表.iter_rows():
            for 单元格 in 行:
                val = 单元格.value
                if val is None:
                    continue
                if isinstance(val, (int, float)):
                    if isinstance(val, float) and val == int(val) and abs(val) >= 1e15:
                        文本 = f"{int(val)}"
                    else:
                        continue
                else:
                    文本 = str(val)
                还原结果 = 映射.批量还原文本(文本)
                if 还原结果 != 文本:
                    单元格.value = 还原结果
    wb.save(输出路径)
    wb.close()
    return 输出路径


def _ppt分配文本到运行(运行列表: list[tuple], 替换后文本: str):
    if not 运行列表:
        return
    总长度 = sum(len(t) for _, t in 运行列表)
    if 总长度 == 0:
        return
    pos = 0
    剩余 = len(替换后文本)
    for i, (run, 原文) in enumerate(运行列表):
        原文长度 = len(原文)
        if i == len(运行列表) - 1:
            分配长度 = 剩余
        else:
            分配长度 = max(1, int(原文长度 / 总长度 * len(替换后文本)))
            分配长度 = min(分配长度, 剩余)
        run.text = 替换后文本[pos:pos + 分配长度]
        pos += 分配长度
        剩余 -= 分配长度


def _ppt脱敏段落(段落, 映射):
    if not 段落.runs:
        return
    运行列表 = [(run, run.text) for run in 段落.runs if run.text]
    if not 运行列表:
        return
    原文 = ''.join(t for _, t in 运行列表)
    替换后 = 映射.批量替换文本(原文)
    if 替换后 == 原文:
        return
    _ppt分配文本到运行(运行列表, 替换后)


def _ppt还原段落(段落, 映射):
    if not 段落.runs:
        return
    运行列表 = [(run, run.text) for run in 段落.runs if run.text]
    if not 运行列表:
        return
    原文 = ''.join(t for _, t in 运行列表)
    替换后 = 映射.批量还原文本(原文)
    if 替换后 == 原文:
        return
    _ppt分配文本到运行(运行列表, 替换后)


def _脱敏ppt(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    from pptx import Presentation
    from 主程序.文档解析器 import 验证office文件格式
    验证office文件格式(路径)
    shutil.copy2(路径, 输出路径)
    prs = Presentation(输出路径)
    for 幻灯片 in prs.slides:
        for 形状 in 幻灯片.shapes:
            if 形状.has_text_frame:
                for 段落 in 形状.text_frame.paragraphs:
                    _ppt脱敏段落(段落, 映射)
            if 形状.has_table:
                for 行 in 形状.table.rows:
                    for 单元格 in 行.cells:
                        for 段落 in 单元格.text_frame.paragraphs:
                            _ppt脱敏段落(段落, 映射)
    prs.save(输出路径)
    return 输出路径


def _还原ppt(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    from pptx import Presentation
    from 主程序.文档解析器 import 验证office文件格式
    验证office文件格式(路径)
    shutil.copy2(路径, 输出路径)
    prs = Presentation(输出路径)
    for 幻灯片 in prs.slides:
        for 形状 in 幻灯片.shapes:
            if 形状.has_text_frame:
                for 段落 in 形状.text_frame.paragraphs:
                    _ppt还原段落(段落, 映射)
            if 形状.has_table:
                for 行 in 形状.table.rows:
                    for 单元格 in 行.cells:
                        for 段落 in 单元格.text_frame.paragraphs:
                            _ppt还原段落(段落, 映射)
    prs.save(输出路径)
    return 输出路径


def _脱敏pdf(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    import fitz
    import tempfile
    import os
    目标目录 = os.path.dirname(输出路径)
    os.makedirs(目标目录, exist_ok=True)
    fd, 临时路径 = tempfile.mkstemp(suffix=".pdf", dir=目标目录)
    os.close(fd)
    doc = None
    try:
        shutil.copy2(路径, 临时路径)
        doc = fitz.open(临时路径)
        映射项 = sorted(映射.正向映射.items(), key=lambda x: len(x[0]), reverse=True)
        for 页面 in doc:
            文本 = 页面.get_text()
            if not 文本.strip():
                continue
            脱敏文本 = 映射.批量替换文本(文本)
            if 脱敏文本 == 文本:
                continue
            for 原文, 代号 in 映射项:
                if 原文 not in 文本:
                    continue
                搜索结果 = 页面.search_for(原文)
                for 矩形 in 搜索结果:
                    页面.add_redact_annot(矩形, 代号, fontname="china-s", fontsize=8)
            页面.apply_redactions()
        doc.save(临时路径)
        doc.close()
        doc = None
        shutil.move(临时路径, 输出路径)
    except Exception:
        if doc is not None:
            try:
                doc.close()
            except Exception:
                pass
        if os.path.exists(临时路径):
            try:
                os.unlink(临时路径)
            except Exception:
                pass
        raise
    return 输出路径


def _还原pdf(路径: str, 映射: 全局映射表, 输出路径: str) -> str:
    import fitz
    import tempfile
    import os
    目标目录 = os.path.dirname(输出路径)
    os.makedirs(目标目录, exist_ok=True)
    fd, 临时路径 = tempfile.mkstemp(suffix=".pdf", dir=目标目录)
    os.close(fd)
    doc = None
    try:
        shutil.copy2(路径, 临时路径)
        doc = fitz.open(临时路径)
        for 页面 in doc:
            文本 = 页面.get_text()
            还原文本 = 映射.批量还原文本(文本)
            if 还原文本 != 文本:
                for 代号, 原文 in 映射.反向映射.items():
                    if 代号 not in 文本:
                        continue
                    搜索结果 = 页面.search_for(代号)
                    for 矩形 in 搜索结果:
                        页面.add_redact_annot(矩形, 原文, fontname="china-s", fontsize=8)
            页面.apply_redactions()
        doc.save(临时路径)
        doc.close()
        doc = None
        shutil.move(临时路径, 输出路径)
    except Exception:
        if doc is not None:
            try:
                doc.close()
            except Exception:
                pass
        if os.path.exists(临时路径):
            try:
                os.unlink(临时路径)
            except Exception:
                pass
        raise
    return 输出路径
