from __future__ import annotations

from pathlib import Path
from 主程序.Excel读取 import Excel公式读取会话, 读取Excel公式值


SUPPORTED_EXTENSIONS: set[str] = {
    ".txt", ".md", ".docx",
    ".xlsx", ".pptx", ".pdf",
}

_OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"


def 检测文件真实格式(文件路径: str) -> str | None:
    try:
        with open(文件路径, "rb") as f:
            头部 = f.read(8)
    except OSError:
        return None
    if 头部.startswith(_OLE2_MAGIC):
        return "ole2"
    if 头部.startswith(_ZIP_MAGIC):
        return "zip"
    return None


def 验证office文件格式(文件路径: str) -> None:
    后缀 = Path(文件路径).suffix.lower()
    if 后缀 not in {".docx", ".xlsx", ".pptx"}:
        return
    真实格式 = 检测文件真实格式(文件路径)
    格式映射 = {".docx": "Word", ".xlsx": "Excel", ".pptx": "PowerPoint"}
    旧版映射 = {".docx": ".doc", ".xlsx": ".xls", ".pptx": ".ppt"}
    if 真实格式 == "ole2":
        raise ValueError(
            f"该文件实际上是旧版{格式映射[后缀]}（{旧版映射[后缀]}）格式，"
            f"仅被改了扩展名为{后缀}。请先用Office另存为真正的{后缀}格式后再处理。"
        )


def _读取文本文件(路径: str) -> str:
    try:
        return Path(路径).read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        return Path(路径).read_text(encoding="gbk", errors="replace")


def 提取文本(文件路径: str) -> str:
    后缀 = Path(文件路径).suffix.lower()
    if 后缀 in {".txt", ".md"}:
        return _读取文本文件(文件路径)
    if 后缀 in {".docx",".xlsx",".pptx",".pdf"}:
        return "\n".join(b.原文 for b in 读取文档块(文件路径))
    if 后缀 == ".docx":
        return _提取docx文本(文件路径)
    if 后缀 == ".doc":
        raise NotImplementedError("旧版 .doc 格式暂不支持，请先转换为 .docx")
    if 后缀 == ".xlsx":
        return _提取excel文本(文件路径)
    if 后缀 == ".xls":
        raise NotImplementedError("旧版 .xls 格式暂不支持，请先转换为 .xlsx")
    if 后缀 == ".pptx":
        return _提取ppt文本(文件路径)
    if 后缀 == ".ppt":
        raise NotImplementedError("旧版 .ppt 格式暂不支持，请先转换为 .pptx")
    if 后缀 == ".pdf":
        return _提取pdf文本(文件路径)
    raise ValueError(f"不支持的文件格式：{后缀}")


def _提取docx文本(路径: str) -> str:
    from docx import Document
    from 主程序._docx_xml工具 import 收集文档全部部件, 提取部件文本
    验证office文件格式(路径)
    doc = Document(路径)
    所有文本: list[str] = []
    for 部件 in 收集文档全部部件(doc):
        所有文本.extend(提取部件文本(部件))
    return "\n".join(所有文本)


def _excel单元格转文本(值) -> str:
    if 值 is None:
        return ""
    if isinstance(值, float):
        if 值 == int(值) and abs(值) >= 1e15:
            return f"{int(值)}"
        if 值 == int(值):
            return str(int(值))
        return str(值)
    return str(值)


def _提取excel文本(路径: str) -> str:
    import re
    from openpyxl import load_workbook
    验证office文件格式(路径)
    wb = load_workbook(路径, read_only=True, data_only=True)
    所有文本: list[str] = []
    _纯数字 = re.compile(r'^[\d,.\s]+$')
    _金额 = re.compile(r'(?:\d{1,3}(?:,\d{3})*|\d+)(?:\.\d{1,2})?\s*(?:万元|亿元|万|亿|元|港元|美元|欧元|日元|英镑)')
    for 工作表 in wb.worksheets:
        for 行 in 工作表.iter_rows(values_only=True):
            行文本 = []
            for 单元格 in 行:
                if 单元格 is None:
                    continue
                if isinstance(单元格, (int, float)):
                    continue
                文本 = str(单元格).strip()
                if not 文本:
                    continue
                if _纯数字.match(文本):
                    continue
                if _金额.search(文本):
                    continue
                行文本.append(文本)
            if 行文本:
                所有文本.append(" ".join(行文本))
    wb.close()
    return "\n".join(所有文本)


def _提取ppt文本(路径: str) -> str:
    from pptx import Presentation
    验证office文件格式(路径)
    prs = Presentation(路径)
    所有文本: list[str] = []
    for 幻灯片 in prs.slides:
        for 形状 in 幻灯片.shapes:
            if 形状.has_text_frame:
                for 段落 in 形状.text_frame.paragraphs:
                    if 段落.text.strip():
                        所有文本.append(段落.text)
            if 形状.has_table:
                for 行 in 形状.table.rows:
                    行文本 = [单元格.text for 单元格 in 行.cells]
                    所有文本.append(" ".join(行文本))
    return "\n".join(所有文本)


def _提取pdf文本(路径: str) -> str:
    import fitz
    doc = fitz.open(路径)
    页面文本列表: list[str] = []
    for 页面 in doc:
        页面文本列表.append(页面.get_text())
    doc.close()
    return "\n".join(页面文本列表)


def 判断是否扫描版pdf(路径: str, 最小文字阈值: int = 50) -> bool:
    import pymupdf
    with pymupdf.open(路径) as doc:
        总字数 = 0
        for page in doc:
            字数 = len(page.get_text().strip())
            总字数 += 字数
            if 字数 < 最小文字阈值:
                for img in page.get_images():
                    if any(r.get_area() > page.rect.get_area()*0.5 for r in page.get_image_rects(img[0])):
                        return True
        return 总字数 == 0


def 收集支持的文件(路径: str, 包含子文件夹: bool = False) -> list[str]:
    目标 = Path(路径)
    if 目标.is_file():
        return [str(目标)] if 目标.suffix.lower() in SUPPORTED_EXTENSIONS else []
    if not 目标.is_dir():
        return []
    迭代器 = 目标.rglob("*") if 包含子文件夹 else 目标.glob("*")
    文件列表 = [
        str(文件)
        for 文件 in 迭代器
        if 文件.is_file() and 文件.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    return sorted(文件列表)


def 遍历Word段落(doc):
    from docx.oxml.ns import qn
    from 主程序._docx_xml工具 import 收集文档全部部件
    for 部件号, 部件 in enumerate(收集文档全部部件(doc)):
        for 段落号, 段落 in enumerate(部件.iter(qn("w:p"))):
            # 嵌套文本框的文字只归属最近段落，避免重复识别和写回。
            节点 = [t for t in 段落.iter(qn("w:t"))
                    if next(t.iterancestors(qn("w:p")), None) is 段落]
            yield f"docx:{部件号}:{段落号}", 节点, 段落


def 遍历PPT段落(prs):
    def 遍历形状(形状列表, 前缀):
        for i, 形状 in enumerate(形状列表):
            键 = f"{前缀}:{i}"
            if hasattr(形状, "shapes"):
                yield from 遍历形状(形状.shapes, 键)
            if 形状.has_text_frame:
                for j, 段落 in enumerate(形状.text_frame.paragraphs):
                    yield f"{键}:p{j}", 段落, ""
            if 形状.has_table:
                for r, 行 in enumerate(形状.table.rows):
                    上下文 = " | ".join(c.text for c in 行.cells)
                    for c, 单元格 in enumerate(行.cells):
                        for j, 段落 in enumerate(单元格.text_frame.paragraphs):
                            yield f"{键}:r{r}:c{c}:p{j}", 段落, 上下文
    for 页号, 幻灯片 in enumerate(prs.slides):
        yield from 遍历形状(幻灯片.shapes, f"pptx:{页号}")


def PPT文字节点(段落):
    from pptx.oxml.ns import qn
    # 保留原来的软换行和域节点，定位只按实际文字节点累计。
    return list(段落._p.iter(qn("a:t")))


def 读取文档块(文件路径: str, Excel会话=None):
    """保留未清洗原文及位置；规则和模型共用同一套位置。"""
    from 主程序.识别结果 import 文本块
    后缀 = Path(文件路径).suffix.lower()
    验证office文件格式(文件路径)
    块 = []
    if 后缀 in {".txt", ".md"}:
        原文 = _读取文本文件(文件路径)
        for i, 行 in enumerate(原文.splitlines(keepends=True)):
            块.append(文本块(f"text:{i}", 行, {"行": i}, 原值=行))
    elif 后缀 == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        for 编号, 节点, 段落 in 遍历Word段落(Document(文件路径)):
            原文 = "".join(t.text or "" for t in 节点)
            表格 = next(段落.iterancestors(qn("w:tbl")), None)
            上下文 = ""
            if 表格 is not None:
                上下文 = " | ".join("".join(t.text or "" for t in tr.iter(qn("w:t")))
                                   for tr in list(表格.findall(qn("w:tr")))[:3])
            if 原文:
                块.append(文本块(编号, 原文, {"段落": 编号}, 上下文, 原文))
    elif 后缀 == ".xlsx":
        from contextlib import ExitStack
        from itertools import groupby
        from openpyxl import load_workbook
        with ExitStack() as 资源:
            wb = load_workbook(文件路径, data_only=False)
            资源.callback(wb.close)
            # openpyxl 普通模式已加载所有存储的格；迭代内部存储避免填满稀疏空矩形。
            表格 = [(s, sorted((c for c in s._cells.values() if c.value is not None),
                              key=lambda c: (c.row, c.column))) for s in wb.worksheets]
            公式位置 = [(s.title,c.coordinate) for s,格 in 表格 for c in 格 if c.data_type == "f"]
            计算值, 计算错误 = {}, ""
            if 公式位置:
                try:
                    计算值 = 读取Excel公式值(文件路径,公式位置,会话=Excel会话)
                except Exception as exc:
                    计算错误 = f"无法计算公式显示值：{exc}"
            缓存 = None
            if any(p not in 计算值 for p in 公式位置):
                缓存 = load_workbook(文件路径, data_only=True)
                资源.callback(缓存.close)
            for i, (sheet,格) in enumerate(表格):
                显示 = {}
                for cell in 格:
                    键 = (sheet.title, cell.coordinate)
                    值 = cell.value
                    if cell.data_type == "f":
                        值 = 计算值[键] if 键 in 计算值 else 缓存[sheet.title][cell.coordinate].value
                    显示[cell.coordinate] = 值
                def 文本(cell):
                    值 = 显示[cell.coordinate]
                    return _excel单元格转文本(cell.value if 值 is None else 值)
                表头 = " | ".join(文本(c) for c in 格 if c.row <= 3)
                for _, 行组 in groupby(格, key=lambda c: c.row):
                    row = list(行组)
                    行上下文 = " | ".join(文本(c) for c in row)
                    for cell in row:
                        位置 = {"工作表": sheet.title, "单元格": cell.coordinate}
                        值 = 显示[cell.coordinate]
                        if cell.data_type == "f":
                            if (sheet.title,cell.coordinate) not in 计算值:
                                位置["公式复算未完成"] = 计算错误 or "公式含不可用引用或计算错误"
                            if 值 is None:
                                位置["公式结果未读取"] = True
                                值 = cell.value
                        原文 = _excel单元格转文本(值)
                        块.append(文本块(f"xlsx:{i}:{cell.coordinate}", 原文, 位置,
                                         f"工作表：{sheet.title}；表头：{表头}；本行：{行上下文}", cell.value, cell.data_type))
    elif 后缀 == ".pptx":
        from pptx import Presentation
        for 编号, 段落, 上下文 in 遍历PPT段落(Presentation(文件路径)):
            原文 = "".join(x.text or "" for x in PPT文字节点(段落))
            if 原文:
                块.append(文本块(编号, 原文, {"段落": 编号}, 上下文, 原文))
    elif 后缀 == ".pdf":
        import pymupdf
        with pymupdf.open(文件路径) as doc:
            for 页号, page in enumerate(doc):
                for b, 内容 in enumerate(page.get_text("rawdict")["blocks"]):
                    for l, 行 in enumerate(内容.get("lines", [])):
                        字符 = [c for span in 行["spans"] for c in span["chars"]]
                        原文 = "".join(c["c"] for c in 字符)
                        if 原文:
                            块.append(文本块(f"pdf:{页号}:{b}:{l}", 原文,
                                {"页": 页号, "字符框": [c["bbox"] for c in 字符], "行框": 行["bbox"],
                                 "字号": min(s["size"] for s in 行["spans"])}, 原值=原文))
    else:
        raise ValueError(f"不支持的文件格式：{后缀}")
    return 块
