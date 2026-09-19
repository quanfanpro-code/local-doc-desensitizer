"""把核准的具体出现写回文件，并保存被改位置的原始值。"""
from __future__ import annotations
from datetime import datetime
import hashlib
from pathlib import Path
import re
import shutil
from 主程序.识别结果 import 按位置替换, 合并出现


def 备份已有输出(路径):
    p = Path(路径)
    if not p.exists():
        return
    目录 = Path.home() / "BackUp" / ("local-doc-desensitizer_输出_" + datetime.now().strftime("%Y%m%d_%H%M%S_%f"))
    目录.mkdir(parents=True)
    副本 = 目录 / p.name
    shutil.copy2(p, 副本)
    if hashlib.sha256(p.read_bytes()).digest() != hashlib.sha256(副本.read_bytes()).digest():
        raise OSError("原有输出文件备份校验失败")


def 准备替换(块列表, 结果, 映射):
    from 主程序.ner引擎 import 判断机构子类型
    for b in 块列表:
        映射.保留原文代号(b.原文)
    for 编号, 机构 in 结果.机构.items():
        全称 = 机构.get("canonical")
        if 全称:
            映射.注册机构称呼(编号, 全称, 判断机构子类型(全称), True)
    命中, 错误 = 合并出现(块列表, 结果.出现)
    结果.未完成.extend(错误)
    替换 = {}
    for m in 命中:
        if m.类型 == "Ni":
            机构 = 结果.机构[m.机构编号]
            全称 = 机构.get("canonical")
            首称 = 机构.get("names", [m.原词])[0]
            代号 = 映射.注册机构称呼(m.机构编号, m.原词, 判断机构子类型(全称 or m.原词),
                                    m.原词 == (全称 or 首称))
        elif m.替换值 is not None:
            if any(m.替换值 in b.原文 for b in 块列表):
                代号 = 映射.查找或创建(m.原词,m.类型)
            else:
                映射.注册自定义映射(m.原词,m.替换值)
                代号 = 映射.正向映射[m.原词]
        else:
            代号 = 映射.查找或创建(m.原词, m.类型)
        替换.setdefault(m.块编号, []).append((m.开始, m.结束, 代号))
    return 替换


def 残留对应原文(原块, 输出块, 替换, 出现):
    """仅在可确定对应的未修改文字区域内，把残留位置映回原文。"""
    from dataclasses import replace
    源表 = {b.编号:b for b in 原块}
    输出表 = {b.编号:b for b in 输出块}
    out = 输出表[出现.块编号]
    src = 源表.get(出现.块编号)
    if src and 按位置替换(src.原文,替换.get(src.编号,[])) == out.原文:
        前原 = 前新 = 0
        for s,e,token in sorted(替换.get(src.编号,[])) + [(len(src.原文),len(src.原文),"")]:
            尾新 = 前新+s-前原
            if 前新 <= 出现.开始 < 出现.结束 <= 尾新:
                return replace(出现,块编号=src.编号,开始=前原+出现.开始-前新,结束=前原+出现.结束-前新)
            前原,前新 = e,尾新+len(token)
    if "页" in out.位置:
        import pymupdf
        目标框 = pymupdf.Rect(out.位置["字符框"][出现.开始])
        for r in out.位置["字符框"][出现.开始+1:出现.结束]:
            目标框 |= pymupdf.Rect(r)
        候选 = []
        for b in 原块:
            if b.位置.get("页") != out.位置["页"]:
                continue
            for m in re.finditer(re.escape(出现.原词),b.原文):
                框 = pymupdf.Rect(b.位置["字符框"][m.start()])
                for r in b.位置["字符框"][m.start()+1:m.end()]:
                    框 |= pymupdf.Rect(r)
                if (框 & 目标框).get_area() > min(框.get_area(),目标框.get_area())*0.5:
                    候选.append(replace(出现,块编号=b.编号,开始=m.start(),结束=m.end()))
        if len(候选) == 1:
            return 候选[0]
    return None


def 替换节点(节点, 替换):
    原字 = [x.text or "" for x in 节点]
    # 先验证所有跨度，再开始修改，避免半段写回后才发现越界。
    按位置替换("".join(原字), 替换)
    偏移 = []
    长度 = 0
    for t in 原字:
        偏移.append(长度); 长度 += len(t)
    for s,e,代号 in sorted(替换, reverse=True):
        首 = next(i for i,t in enumerate(原字) if 偏移[i] <= s < 偏移[i]+len(t))
        尾 = next(i for i,t in enumerate(原字) if 偏移[i] < e <= 偏移[i]+len(t))
        a,b = s-偏移[首], e-偏移[尾]
        if 首 == 尾:
            节点[首].text = (节点[首].text or "")[:a] + 代号 + (节点[首].text or "")[b:]
        else:
            节点[首].text = (节点[首].text or "")[:a] + 代号
            for i in range(首+1, 尾):
                节点[i].text = ""
            节点[尾].text = (节点[尾].text or "")[b:]
    return 原字


def _PDF字符(page):
    return [c for b in page.get_text("rawdict")["blocks"] for line in b.get("lines",[])
            for span in line["spans"] for c in span["chars"]]


def _PDF区域字符框(page, 区域):
    import pymupdf
    rect = pymupdf.Rect(区域)
    return [c["bbox"] for c in _PDF字符(page)
            if rect.contains(pymupdf.Point((c["bbox"][0]+c["bbox"][2])/2,(c["bbox"][1]+c["bbox"][3])/2))]


def _PDF清除字符(page, 字符框):
    """MuPDF 按边界框相交删除整字；用字内小区域触发，先排除邻字碰撞。"""
    import pymupdf
    目标 = {tuple(r) for r in 字符框}
    保留 = [pymupdf.Rect(c["bbox"]) for c in _PDF字符(page) if tuple(c["bbox"]) not in 目标]
    for r in 目标:
        cx,cy = (r[0]+r[2])/2,(r[1]+r[3])/2
        dx,dy = min((r[2]-r[0])/4,0.1),min((r[3]-r[1])/4,0.1)
        框 = pymupdf.Rect(cx-dx,cy-dy,cx+dx,cy+dy)
        if 框.is_empty or any((框 & p).get_area() > 0 for p in 保留):
            raise ValueError(f"PDF目标文字与其他文字边界重叠，无法安全清除，页 {page.number+1}")
        page.add_redact_annot(框,fill=False)


def _PDF插入(page, 矩形, 文字, 字号=10):
    import pymupdf
    字体文件 = Path(r"C:\Windows\Fonts\msyh.ttc")
    参数 = {"fontname":"maskfont","fontfile":str(字体文件)} if 字体文件.exists() else {"fontname":"china-s"}
    for 大小 in sorted({字号,10,8,6,5,4},reverse=True):
        if 大小 > 字号:
            continue
        剩余 = page.insert_textbox(pymupdf.Rect(矩形),文字,fontsize=大小,**参数)
        if 剩余 >= 0:
            return
    raise ValueError(f"PDF位置无法容纳替换文字，页 {page.number+1}")


def 保存(原路径, 块列表, 结果, 映射, 输出路径):
    from 主程序.文档解析器 import 遍历Word段落, 遍历PPT段落, PPT文字节点
    if Path(原路径).resolve() == Path(输出路径).resolve():
        raise ValueError("输出路径不能覆盖原件")
    替换 = 准备替换(块列表,结果,映射)
    后缀 = Path(原路径).suffix.lower()
    备份已有输出(输出路径)
    if 后缀 in {".txt",".md"}:
        内容 = "".join(按位置替换(b.原文,替换.get(b.编号,[])) for b in 块列表)
        Path(输出路径).write_text(内容,encoding="utf-8-sig")
    elif 后缀 in {".docx",".pptx"}:
        if 后缀 == ".docx":
            from 主程序._docx_xml工具 import 打开Word as Document
            doc = Document(原路径)
            段落 = [(k,节点) for k,节点,p in 遍历Word段落(doc)]
        else:
            from pptx import Presentation
            doc = Presentation(原路径)
            段落 = [(k,PPT文字节点(p)) for k,p,ctx in 遍历PPT段落(doc)]
        原块 = {b.编号:b for b in 块列表}
        for 编号,节点 in 段落:
            spans = 替换.get(编号)
            if not spans:
                continue
            if "".join(x.text or "" for x in 节点) != 原块[编号].原文:
                raise ValueError(f"文档段落结构变化，不能按原位置写回：{编号}")
            原字 = 替换节点(节点,spans)
            映射.文件记录.append({"格式":后缀,"块编号":编号,"原节点文字":原字,
                "输出节点文字":[x.text or "" for x in 节点]})
        doc.save(输出路径)
    elif 后缀 == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.formula import Tokenizer
        wb = load_workbook(原路径)
        单元格块 = {}
        for b in 块列表:
            单元格块.setdefault((b.位置["工作表"],b.位置["单元格"]),[]).append(b)
        处理位置 = {k for k,bs in 单元格块.items() if any(b.编号 in 替换 for b in bs)}
        try:
            for 位置,bs in 单元格块.items():
                if 位置 not in 处理位置:
                    continue
                b = next(x for x in bs if x.数据类型 != "formula_text")
                spans = 替换.get(b.编号,[])
                cell = wb[位置[0]][位置[1]]
                输出值 = 按位置替换(b.原文,spans)
                if b.数据类型 == "f":
                    公式 = Tokenizer(b.原值)
                    tokens = 公式.items
                    单常量 = (len(tokens) == 1 and tokens[0].subtype == "TEXT"
                              and tokens[0].value[1:-1].replace('""','"') == b.原文)
                    for 常量 in bs:
                        修改 = 替换.get(常量.编号)
                        if 常量.数据类型 != "formula_text" or not 修改:
                            continue
                        token = tokens[常量.位置["公式常量序号"]]
                        if token.subtype != "TEXT" or token.value[1:-1].replace('""','"') != 常量.原文:
                            raise ValueError(f"公式常量结构变化，不能按原位置写回：{常量.编号}")
                        token.value = '"' + 按位置替换(常量.原文,修改).replace('"','""') + '"'
                    if not spans:
                        # 隐藏分支、参数里的文字可定位修改，计算和引用仍由原公式负责。
                        输出值 = 公式.render()
                    elif 单常量:
                        输出值 = '="' + 输出值.replace('"','""') + '"'
                    elif len(tokens) == 1 and tokens[0].subtype == "RANGE":
                        m = re.fullmatch(r"(?:(?:'((?:[^']|'')+)'|([^!]+))!)?(\$?[A-Z]+\$?\d+)",tokens[0].value)
                        if m:
                            sheet = (m.group(1) or m.group(2) or b.位置["工作表"]).replace("''","'")
                            坐标 = m.group(3).replace("$","")
                            if (sheet,坐标) in 处理位置:
                                continue
                    # 敏感计算结果无法由单个常量或已处理引用可靠修正时，仅该格转值。
                原值 = cell.value
                日期值 = isinstance(原值,datetime)
                映射.文件记录.append({"格式":后缀,"块编号":b.编号,"位置":b.位置,
                    "原值":原值.isoformat() if 日期值 else 原值,"原值为日期":日期值,
                    "原类型":cell.data_type,"输出值":输出值})
                cell.value = 输出值
                if not (b.数据类型 == "f" and isinstance(输出值,str) and 输出值.startswith("=")):
                    cell.data_type = "s"
            wb.calculation.fullCalcOnLoad = True
            wb.calculation.forceFullCalc = True
            wb.save(输出路径)
        finally:
            wb.close()
    elif 后缀 == ".pdf":
        import pymupdf
        with pymupdf.open(原路径) as doc:
            待写, 清除 = [], {}
            for b in 块列表:
                for s,e,代号 in 替换.get(b.编号,[]):
                    框 = pymupdf.Rect(b.位置["字符框"][s])
                    for rect in b.位置["字符框"][s+1:e]:
                        框 |= pymupdf.Rect(rect)
                    page = doc[b.位置["页"]]
                    清除.setdefault(page.number,[]).extend(b.位置["字符框"][s:e])
                    待写.append({"格式":后缀,"块编号":f"{b.编号}:{s}",
                        "页":page.number,"矩形":list(框),"原文":b.原文[s:e],"输出值":代号,
                        "字号":b.位置.get("字号",10)})
            for page in doc:
                _PDF清除字符(page,清除.get(page.number,[]))
                # 保留背景图和线条；整页扫描图已在入口转为 Markdown 输出。
                page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,graphics=0)
            for r in 待写:
                _PDF插入(doc[r["页"]],r["矩形"],r["输出值"],r["字号"])
            doc.save(输出路径,garbage=4,deflate=True)
            映射.文件记录.extend(待写)
    else:
        raise ValueError(f"尚未接入该格式的定位写回：{后缀}")
    return 输出路径


def 还原(文件路径, 映射, 输出路径):
    from 主程序.文档解析器 import 遍历Word段落, 遍历PPT段落, PPT文字节点
    后缀 = Path(文件路径).suffix.lower()
    记录 = {r["块编号"]:r for r in 映射.文件记录 if r["格式"] == 后缀}
    if 后缀 == ".pdf":
        import pymupdf
        with pymupdf.open(文件路径) as doc:
            for r in 记录.values():
                page = doc[r["页"]]
                区域文字 = "".join(page.get_textbox(pymupdf.Rect(r["矩形"])).split())
                if "".join(r["输出值"].split()) not in 区域文字:
                    raise ValueError("PDF中的代号与还原映射不一致")
                _PDF清除字符(page,_PDF区域字符框(page,r["矩形"]))
            for page in doc:
                page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE,graphics=0)
            for r in 记录.values():
                _PDF插入(doc[r["页"]],r["矩形"],r["原文"],r["字号"])
            备份已有输出(输出路径)
            doc.save(输出路径,garbage=4,deflate=True)
        return 输出路径
    if 后缀 == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(文件路径)
        try:
            for r in 记录.values():
                cell = wb[r["位置"]["工作表"]][r["位置"]["单元格"]]
                if cell.value != r["输出值"]:
                    raise ValueError(f"脱敏后单元格被修改，无法准确还原：{r['块编号']}")
                cell.value = datetime.fromisoformat(r["原值"]) if r.get("原值为日期") else r["原值"]
                cell.data_type = r["原类型"]
            备份已有输出(输出路径)
            wb.save(输出路径)
        finally:
            wb.close()
        return 输出路径
    if 后缀 == ".docx":
        from 主程序._docx_xml工具 import 打开Word as Document
        doc = Document(文件路径)
        段落 = [(k,节点) for k,节点,p in 遍历Word段落(doc)]
    elif 后缀 == ".pptx":
        from pptx import Presentation
        doc = Presentation(文件路径)
        段落 = [(k,PPT文字节点(p)) for k,p,ctx in 遍历PPT段落(doc)]
    else:
        raise ValueError(f"尚未接入该格式的定位还原：{后缀}")
    已恢复 = set()
    for 编号,节点 in 段落:
        if 编号 not in 记录:
            continue
        r = 记录[编号]
        if [x.text or "" for x in 节点] != r["输出节点文字"]:
            raise ValueError(f"脱敏后段落被修改，无法准确还原：{编号}")
        for node,原文 in zip(节点,r["原节点文字"]):
            node.text = 原文
        已恢复.add(编号)
    if 已恢复 != set(记录):
        raise ValueError("还原映射中的部分位置在文件中不存在")
    备份已有输出(输出路径)
    doc.save(输出路径)
    return 输出路径
