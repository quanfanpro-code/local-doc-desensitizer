"""验证公式批量传输、资源释放与实际脱敏入口；材料保留供复核。"""
from pathlib import Path
from types import SimpleNamespace
import hashlib
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter
from 主程序 import 文档解析器
from 主程序.识别结果 import 识别结果
from 维护工具.运行脱敏验证 import 验证用临时目录


class Excel读取测试(unittest.TestCase):
    def 会话(self, **kwargs):
        类 = getattr(文档解析器, "Excel公式读取会话", None)
        self.assertIsNotNone(类, "缺少任务内可复用的 Excel 公式读取会话")
        return 类(**kwargs)

    def 假Excel(self, 数值):
        区域 = []
        def 范围(address):
            c1,r1,c2,r2 = range_boundaries(address)
            区域.append((r2-r1+1)*(c2-c1+1))
            self.assertLessEqual(区域[-1],4096, "稀疏坐标不得扩成巨大读取区域")
            值 = tuple(tuple(数值.get(f"{get_column_letter(c)}{r}") for c in range(c1,c2+1)) for r in range(r1,r2+1))
            return SimpleNamespace(Value2=值[0][0] if 区域[-1] == 1 else 值)
        sheet = SimpleNamespace(Range=范围, Evaluate=lambda expr: expr == "ISERROR(A4)")
        book = Mock(); book.Worksheets.side_effect = lambda name: sheet
        app = Mock(); app.Workbooks.Open.return_value = book
        return app,book,区域

    def test_有限批量取值及单格稀疏边界(self):
        app,book,区域 = self.假Excel({**{f"A{i}":"星河" for i in range(1,5001)},"XFD1048576":"远景"})
        with patch("win32com.client.DispatchEx",return_value=app):
            with self.会话() as 会话:
                值 = 会话.读取("样例.xlsx",[("数据",f"A{i}") for i in range(1,5001)]+[("数据","XFD1048576")])
        self.assertEqual(len(值),5001)
        self.assertEqual(值[("数据","A5000")],"星河")
        self.assertEqual(值[("数据","XFD1048576")],"远景")
        self.assertLess(len(区域),10, "不能退回逐格跨进程读取")
        self.assertEqual(book.Close.call_args.kwargs,{"SaveChanges":False})
        self.assertEqual(app.Workbooks.Open.call_args.kwargs,{"UpdateLinks":0,"ReadOnly":True})
        self.assertEqual(app.AutomationSecurity,3)
        self.assertFalse(app.EnableEvents)
        book.Close.assert_called_once_with(SaveChanges=False)
        app.Quit.assert_called_once()

    def test_零空文本布尔值与错误码同值的数字(self):
        app,book,区域 = self.假Excel({"A1":0,"A2":False,"A3":"","A4":-2146826281,"A5":-2146826281})
        with patch("win32com.client.DispatchEx",return_value=app):
            with self.会话() as 会话:
                值 = 会话.读取("样例.xlsx",[("数据",f"A{i}") for i in range(1,6)])
        self.assertEqual(值,{("数据","A1"):0,("数据","A2"):False,("数据","A3"):"",("数据","A5"):-2146826281})

    def test_取消与读取异常均释放自己的资源(self):
        for 取消 in (False,True):
            with self.subTest(取消=取消):
                app,book,_ = self.假Excel({"A1":"星河"})
                状态 = [False]
                if 取消:
                    app.CalculateFullRebuild.side_effect = lambda: 状态.__setitem__(0,True)
                else:
                    app.CalculateFullRebuild.side_effect = RuntimeError("故障注入")
                with patch("win32com.client.DispatchEx",return_value=app),patch("pythoncom.CoUninitialize") as 释放:
                    with self.assertRaises(RuntimeError):
                        with self.会话(取消检查=lambda:状态[0]) as 会话:
                            会话.读取("样例.xlsx",[("数据","A1")])
                book.Close.assert_called_once_with(SaveChanges=False)
                app.Quit.assert_called_once(); 释放.assert_called_once()

    def test_关闭失败不吞掉交付结果且仍释放COM(self):
        app,book,_ = self.假Excel({"A1":"星河"})
        app.Quit.side_effect = RuntimeError("关闭失败")
        with patch("win32com.client.DispatchEx",return_value=app),patch("pythoncom.CoUninitialize") as 释放:
            with self.会话() as 会话:
                值 = 会话.读取("样例.xlsx",[("数据","A1")])
        self.assertEqual(值[("数据","A1")],"星河")
        self.assertTrue(会话.问题); 释放.assert_called_once()

    def test_单个文件打不开不妨碍下一份文件(self):
        app,book,_ = self.假Excel({"A1":"星河"})
        app.Workbooks.Open.side_effect = [RuntimeError("文件打不开"),book]
        with patch("win32com.client.DispatchEx",return_value=app):
            with self.会话() as 会话:
                with self.assertRaisesRegex(RuntimeError,"文件打不开"):
                    会话.读取("坏文件.xlsx",[("数据","A1")])
                self.assertEqual(会话.读取("好文件.xlsx",[("数据","A1")]),{("数据","A1"):"星河"})
        self.assertFalse(会话.问题)

    def test_工作簿关闭失败不继续复用并报告(self):
        app,book,_ = self.假Excel({"A1":"星河"})
        book.Close.side_effect = RuntimeError("关闭失败")
        with patch("win32com.client.DispatchEx",return_value=app):
            with self.会话() as 会话:
                self.assertEqual(会话.读取("样例.xlsx",[("数据","A1")]),{("数据","A1"):"星河"})
                with self.assertRaisesRegex(RuntimeError,"工作簿关闭失败"):
                    会话.读取("下一份.xlsx",[("数据","A1")])
        self.assertTrue(会话.问题); app.Quit.assert_called_once()

    def test_无公式不启动Excel且空读取不初始化(self):
        with patch("win32com.client.DispatchEx",side_effect=AssertionError("不应启动Excel")):
            with self.会话() as 会话:
                self.assertEqual(会话.读取("不存在.xlsx",[]),{})
                with 验证用临时目录() as d:
                    p = Path(d)/"普通.xlsx"; wb = Workbook(); wb.active["A1"]="星河"; wb.save(p); wb.close()
                    self.assertEqual(文档解析器.读取文档块(str(p),Excel会话=会话)[0].原文,"星河")

    def test_稀疏表只解析有值位置(self):
        with 验证用临时目录() as d:
            p=Path(d)/"稀疏.xlsx"; wb=Workbook(); wb.active["A1"]="供应商"; wb.active["XFD1048576"]="星河"; wb.save(p); wb.close()
            # 防止旧实现真的分配整张 Excel 的空白矩形。
            with patch("openpyxl.worksheet.worksheet.Worksheet.iter_rows",side_effect=AssertionError("不得填满稀疏矩形")):
                块=文档解析器.读取文档块(str(p))
            self.assertEqual([b.原文 for b in 块],["供应商","星河"])

    def test_缺Excel不重复启动且旧缓存标为未完成(self):
        from zipfile import ZipFile
        with 验证用临时目录() as d:
            初稿=Path(d)/"初稿.xlsx"; p=Path(d)/"缓存.xlsx"
            wb=Workbook(); wb.active["A1"]='="新机构"'; wb.save(初稿); wb.close()
            with ZipFile(初稿) as src,ZipFile(p,"w") as dst:
                for name in src.namelist():
                    data=src.read(name)
                    if name=="xl/worksheets/sheet1.xml":
                        data=data.replace(b'<c r="A1">',b'<c r="A1" t="str">').replace(b'<v></v>',"<v>旧机构</v>".encode())
                    dst.writestr(name,data)
            with patch("win32com.client.DispatchEx",side_effect=RuntimeError("Excel不可用")) as 启动:
                with self.会话() as 会话:
                    for _ in range(2):
                        块=文档解析器.读取文档块(str(p),Excel会话=会话)
                        self.assertEqual(块[0].原文,"旧机构")
                        self.assertIn("公式复算未完成",块[0].位置)
            self.assertEqual(启动.call_count,1)
            块=文档解析器.读取文档块(str(p))
            self.assertEqual(块[0].原文,"新机构")
            self.assertNotIn("公式复算未完成",块[0].位置)

    def test_输出复算失败保留文件但不报告完全成功(self):
        from 主程序.脱敏处理器 import 脱敏处理器
        from 主程序.识别结果 import 文本块
        with 验证用临时目录() as d:
            p=Path(d)/"原件.xlsx"; wb=Workbook(); wb.active["A1"]="=SUM(1,2)"; wb.save(p); wb.close()
            原=文本块("xlsx:0:A1","3",{"工作表":"Sheet","单元格":"A1"},原值="=SUM(1,2)",数据类型="f")
            出=文本块("xlsx:0:A1","3",{"工作表":"Sheet","单元格":"A1","公式复算未完成":"输出计算失败"},原值="=SUM(1,2)",数据类型="f")
            proc=脱敏处理器()
            with patch.object(文档解析器,"读取文档块",side_effect=[[原],[出]]),patch.object(proc._ner,"识别文档",return_value=识别结果()),patch.object(proc._ner,"复查输出",return_value=识别结果()):
                结果=proc.处理文件列表([str(p)],d)
            self.assertEqual(结果.成功数,0)
            self.assertEqual(len(结果.已生成文件),1)
            self.assertTrue(any("输出计算失败" in str(x) for x in 结果.未完成环节))

    def test_真实Excel计算混合值及上下文并保留原件(self):
        with 验证用临时目录() as d:
            p = Path(d)/"混合.xlsx"; wb=Workbook(); ws=wb.active; ws.title="数据"
            for row in [['="供应商"','=0','=FALSE()','=""','=1/0','=-2146826281'],['="星"&"河"',"='数据'!A2"]]:
                ws.append(row)
            wb.save(p); wb.close(); 原件=hashlib.sha256(p.read_bytes()).hexdigest()
            with self.会话() as 会话:
                块=文档解析器.读取文档块(str(p),Excel会话=会话)
            by={b.位置["单元格"]:b for b in 块}
            self.assertEqual([by[c].原文 for c in ("A1","B1","C1","D1","A2","B2","F1")],["供应商","0","False","","星河","星河","-2146826281"])
            self.assertIn("公式复算未完成",by["E1"].位置)
            self.assertNotIn("公式复算未完成",by["D1"].位置)
            self.assertIn("本行：星河 | 星河",by["A2"].上下文)
            self.assertEqual(by["A2"].原值,'="星"&"河"')
            self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),原件)

    def test_真实批处理复查共用实例且可还原公式(self):
        import win32com.client
        from 主程序.脱敏处理器 import 脱敏处理器
        with 验证用临时目录() as d:
            文件=[]; 原件=[]
            for 名称 in ("甲","乙"):
                p=Path(d)/f"{名称}.xlsx"; wb=Workbook(); ws=wb.active
                ws["A1"]='="138"&"00138000"'; ws["B1"]="=SUM(1,2)"; wb.save(p); wb.close()
                文件.append(str(p)); 原件.append(p.read_bytes())
            处理器=脱敏处理器()
            with patch("win32com.client.DispatchEx",wraps=win32com.client.DispatchEx) as 启动,patch.object(处理器._ner,"识别文档",side_effect=lambda *a:识别结果()),patch.object(处理器._ner,"复查输出",side_effect=lambda *a:识别结果()):
                结果=处理器.处理文件列表(文件,d)
            self.assertEqual(结果.成功数,2,结果.未完成环节)
            self.assertEqual(启动.call_count,1,"原件及输出复查应复用本任务实例")
            for p in 结果.已生成文件:
                w=load_workbook(p); self.assertEqual(w.active["A1"].value,"[手机号1]"); self.assertEqual(w.active["B1"].value,"=SUM(1,2)"); w.close()
            还原=处理器.还原文件列表(结果.已生成文件,输出目录=d)
            self.assertEqual(还原.成功数,2)
            for p in 还原.成功文件:
                w=load_workbook(p); self.assertEqual(w.active["A1"].value,'="138"&"00138000"'); w.close()
            self.assertEqual([Path(p).read_bytes() for p in 文件],原件)


if __name__ == "__main__":
    unittest.main()
