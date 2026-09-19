"""独立复核失败样例的行为回归，材料由正式验证入口保留。"""
from pathlib import Path
import unittest
from unittest.mock import patch

from 维护工具.运行脱敏验证 import 验证用临时目录
from 主程序.固定规则 import 固定规则识别
from 主程序.识别结果 import 文本块, 出现记录, 识别结果
from 主程序.脱敏处理器 import 脱敏处理器


class 固定与合并修复测试(unittest.TestCase):
    def 检出(self, 原文, 字段="", 金额=False, 日期=False, 月日=False):
        位置 = {"列字段":字段} if 字段 else {}
        return 固定规则识别(文本块("p",原文,位置),日期,月日,金额)

    def test_明确短账号完整处理但普通编号不误改(self):
        for 号码 in ("1234567890","123456789012","1234567890123456789012"):
            for 原文,字段 in (("收款银行账号："+号码,""),(号码,"银行账号")):
                with self.subTest(原文=原文,字段=字段):
                    self.assertEqual([(m.原词,m.类型) for m in self.检出(原文,字段)],[(号码,"bank_account")])
        for 原文 in ("产品编号：123456789012","数量：1234567890","[银行账号1]"):
            self.assertFalse(self.检出(原文))

    def test_金额字段纯数字沿用币种门槛而数量不误改(self):
        for 原文,字段,预期 in (("1234567.89","金额（元）","100万余元"),
                ("2000000","金额（元）","200万余元"),
                ("123","金额（万元）","100万余元"),
                ("合同金额（美元）：1234567.89","","100万余美元")):
            with self.subTest(原文=原文,字段=字段):
                self.assertEqual([m.替换值 for m in self.检出(原文,字段,金额=True) if m.类型=="amount"],[预期])
        self.assertFalse(self.检出("1234567", "数量", 金额=True))
        self.assertFalse(self.检出("合同金额（元）：1234567.89"))

    def test_MAC不吞标签并支持中文和同种分隔符(self):
        for 原文,号码 in (("MAC：00:11:22:33:44:55。","00:11:22:33:44:55"),
                ("设备00:11:22:33:44:55已连接","00:11:22:33:44:55"),
                ("网卡AA-BB-CC-DD-EE-FF","AA-BB-CC-DD-EE-FF"),
                ("网卡0012.3456.789A","0012.3456.789A")):
            with self.subTest(原文=原文):
                self.assertEqual([m.原词 for m in self.检出(原文) if m.类型=="mac_address"],[号码])
        for 原文 in ("MAC：00:11-22:33:44:55","型号X00:11:22:33:44:55Z"):
            self.assertFalse([m for m in self.检出(原文) if m.类型=="mac_address"])

    def test_车牌中文相邻小写分隔点及既有特殊尾字(self):
        for 原文,号码 in (("车辆川A12345已入场。","川A12345"),
                ("车辆川a12345已入场。","川a12345"),
                ("车牌川A·23047B最近10分钟","川A·23047B"),
                ("车辆川A1234学已入场","川A1234学")):
            with self.subTest(原文=原文):
                self.assertEqual([m.原词 for m in self.检出(原文) if m.类型=="plate_number"],[号码])
        for 原文 in ("型号X川A12345Y","川A123456789","产品AB12345"):
            self.assertFalse([m for m in self.检出(原文) if m.类型=="plate_number"])

    def test_分隔身份证保留完整原文跨度(self):
        for 号码 in ("110105 19491231 002X","110105-19491231-002X","１１０１０５ １９４９１２３１ ００２ｘ"):
            with self.subTest(号码=号码):
                self.assertEqual([(m.原词,m.类型) for m in self.检出("身份证："+号码)],[(号码,"id_card")])

    def test_无效日期不因月日开关退回局部命中(self):
        for 原文 in ("2026年99月99日","2026年2月29日","2026-04-31","二〇二六年二月二十九日"):
            with self.subTest(原文=原文):
                self.assertFalse([m for m in self.检出(原文,日期=True,月日=True) if m.类型=="date"])
        for 原文 in ("2024年2月29日","二〇二四年二月二十九日","2026-09","2月29日"):
            self.assertTrue([m for m in self.检出(原文,日期=True,月日=True) if m.类型=="date"])

    def test_模型错误跨度不能丢弃完整手机号并可还原(self):
        with 验证用临时目录() as d:
            p=Path(d)/"模型冲突.txt";p.write_text("联系人13800138000。",encoding="utf-8-sig")
            proc=脱敏处理器(启用日期=False)
            模型=识别结果(出现=[出现记录("text:0",0,9,"联系人138001","Nh")])
            with patch.object(proc._ner,"识别文档",return_value=模型),patch.object(proc._ner,"复查输出",return_value=识别结果()):
                result=proc.处理文件列表([str(p)],str(Path(d)/"输出"))
            self.assertEqual(Path(result.已生成文件[0]).read_text(encoding="utf-8-sig"),"联系人[手机号1]。")
            self.assertTrue(any("冲突" in r.get("环节","") for r in result.未完成环节))
            restored=proc.还原文件列表(result.已生成文件,输出目录=str(Path(d)/"还原"))
            self.assertEqual(Path(restored.已生成文件[0]).read_text(encoding="utf-8-sig"),"联系人13800138000。")


class Office修复测试(unittest.TestCase):
    def test_Word脚注尾注写回和重载还原(self):
        import hashlib,zipfile
        from docx import Document
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        with 验证用临时目录() as d:
            p=Path(d)/"注释.docx";doc=Document();para=doc.add_paragraph("正文保持")
            for name,tag,num,phone in (("footnotes","footnote",1,"13800138000"),("endnotes","endnote",2,"13900139000")):
                blob=(f'<w:{name} xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                      f'<w:{tag} w:id="{num}"><w:p><w:r><w:t>联系{phone}</w:t></w:r></w:p></w:{tag}></w:{name}>').encode()
                part=Part(PackURI(f"/word/{name}.xml"),f"application/vnd.openxmlformats-officedocument.wordprocessingml.{name}+xml",blob,doc.part.package)
                doc.part.relate_to(part,f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{name}")
                ref=OxmlElement(f"w:{tag}Reference");ref.set(qn("w:id"),str(num));para.add_run()._r.append(ref)
            doc.save(p);before=hashlib.sha256(p.read_bytes()).digest()
            proc=脱敏处理器(跳过模型=True,启用日期=False)
            result=proc.处理文件列表([str(p)],str(Path(d)/"输出"))
            self.assertEqual(result.成功数,1,(result.失败文件,result.未完成环节))
            out=result.已生成文件[0]
            with zipfile.ZipFile(out) as z:
                self.assertNotIn("13800138000",z.read("word/footnotes.xml").decode())
                self.assertNotIn("13900139000",z.read("word/endnotes.xml").decode())
            restored=proc.还原文件列表([out],输出目录=str(Path(d)/"还原"))
            self.assertEqual(restored.成功数,1,restored.失败文件)
            with zipfile.ZipFile(restored.已生成文件[0]) as z:
                self.assertIn("13800138000",z.read("word/footnotes.xml").decode())
                self.assertIn("13900139000",z.read("word/endnotes.xml").decode())
            self.assertEqual(Document(out).paragraphs[0].text,"正文保持")
            self.assertEqual(hashlib.sha256(p.read_bytes()).digest(),before)

    def test_Excel隐藏常量处理且普通公式引用和原类型可还原(self):
        import hashlib
        from openpyxl import Workbook,load_workbook
        with 验证用临时目录() as d:
            p=Path(d)/"公式.xlsx";wb=Workbook();ws=wb.active;ws.title="原表"
            originals={
                "A1":'=IF(TRUE,"公开信息","13800138000")',
                "B1":'=LEN("13900139000")',
                "C1":"=SUM(1,2)",
                "D1":"='原表'!C1",
                "A2":'=IF(TRUE,"公开""引号","13800138000")',
            }
            for cell,value in originals.items():ws[cell]=value
            amounts=wb.create_sheet("金额");amounts.append(["金额（元）","数量"])
            amounts.append([1234567.89,1234567]);amounts.append([2000000,2000000])
            for coord in ("A2","A3"):amounts[coord].number_format="#,##0.00"
            wb.save(p);wb.close();before=hashlib.sha256(p.read_bytes()).digest()
            proc=脱敏处理器(跳过模型=True,启用日期=False,启用金额=True)
            result=proc.处理文件列表([str(p)],str(Path(d)/"输出"))
            self.assertEqual(result.成功数,1,(result.失败文件,result.未完成环节))
            out=result.已生成文件[0];book=load_workbook(out)
            self.assertEqual(book["原表"]["A1"].value,'=IF(TRUE,"公开信息","[手机号1]")')
            self.assertEqual(book["原表"]["B1"].value,'=LEN("[手机号2]")')
            self.assertEqual(book["原表"]["A2"].value,'=IF(TRUE,"公开""引号","[手机号1]")')
            self.assertEqual(book["原表"]["C1"].value,"=SUM(1,2)")
            self.assertEqual(book["原表"]["D1"].value,"='原表'!C1")
            self.assertEqual(book["金额"]["A2"].value,"100万余元")
            self.assertEqual(book["金额"]["A3"].value,"200万余元")
            self.assertEqual(book["金额"]["B2"].value,1234567)
            book.close()
            restored=proc.还原文件列表([out],输出目录=str(Path(d)/"还原"))
            self.assertEqual(restored.成功数,1,restored.失败文件)
            book=load_workbook(restored.已生成文件[0])
            self.assertEqual({cell:book["原表"][cell].value for cell in originals},originals)
            self.assertEqual(book["金额"]["A2"].value,1234567.89)
            self.assertIsInstance(book["金额"]["A3"].value,int)
            book.close();self.assertEqual(hashlib.sha256(p.read_bytes()).digest(),before)


class PDF修复测试(unittest.TestCase):
    def test_邻行文字在脱敏及还原后均完整(self):
        import hashlib,pymupdf
        with 验证用临时目录() as d:
            p=Path(d)/"邻行.pdf"
            with pymupdf.open() as doc:
                page=doc.new_page()
                page.insert_text((72,72),"13800138000",fontsize=12)
                page.insert_text((72,84),"Ordinary text stays",fontsize=12)
                doc.save(p)
                page.get_pixmap(matrix=pymupdf.Matrix(2,2)).save(str(Path(d)/"原件.png"))
            before=hashlib.sha256(p.read_bytes()).digest()
            proc=脱敏处理器(跳过模型=True,启用日期=False)
            result=proc.处理文件列表([str(p)],str(Path(d)/"输出"))
            self.assertEqual(result.成功数,1,(result.失败文件,result.未完成环节))
            out=result.已生成文件[0]
            with pymupdf.open(out) as doc:
                text=doc[0].get_text()
                self.assertIn("Ordinary text stays",text)
                self.assertNotIn("13800138000",text)
                doc[0].get_pixmap(matrix=pymupdf.Matrix(2,2)).save(str(Path(d)/"脱敏.png"))
            restored=proc.还原文件列表([out],输出目录=str(Path(d)/"还原"))
            self.assertEqual(restored.成功数,1,restored.失败文件)
            with pymupdf.open(restored.已生成文件[0]) as doc:
                self.assertIn("Ordinary text stays",doc[0].get_text())
                self.assertIn("13800138000",doc[0].get_text())
                doc[0].get_pixmap(matrix=pymupdf.Matrix(2,2)).save(str(Path(d)/"还原.png"))
            self.assertEqual(hashlib.sha256(p.read_bytes()).digest(),before)

    def test_已有OCR层扫描件直接脱敏MD不交付残留原图PDF(self):
        import pymupdf
        from types import SimpleNamespace
        from 主程序.文档解析器 import 判断是否扫描版pdf
        with 验证用临时目录() as d:
            text="Contact 13800138000. This is a searchable scan with visible page pixels and OCR text."
            with pymupdf.open() as image:
                page=image.new_page();page.insert_text((72,72),text,fontsize=11)
                png=page.get_pixmap().tobytes("png")
            p=Path(d)/"已有OCR.pdf"
            with pymupdf.open() as doc:
                page=doc.new_page();page.insert_image(page.rect,stream=png)
                page.insert_text((72,72),text,fontsize=11,render_mode=3)
                doc.save(p)
            self.assertTrue(判断是否扫描版pdf(str(p)))
            ocr=SimpleNamespace(是否可用=lambda:False,pdf转文本=lambda *args: self.fail("已有可用文字层不应重新OCR"))
            proc=脱敏处理器(ocr引擎=ocr,跳过模型=True,启用日期=False)
            result=proc.处理文件列表([str(p)],str(Path(d)/"输出"))
            self.assertEqual(result.成功数,1,(result.失败文件,result.未完成环节))
            self.assertEqual(Path(result.已生成文件[0]).suffix,".md")
            self.assertFalse(list((Path(d)/"输出").glob("*.pdf")))
            output=Path(result.已生成文件[0]).read_text(encoding="utf-8-sig")
            self.assertNotIn("13800138000",output)
            self.assertIn("[手机号1]",output)

    def test_实际输出缺失未命中文字不能报告完成(self):
        from 主程序.ner引擎 import 全局映射表
        proc=脱敏处理器(跳过模型=True,启用日期=False)
        original=[文本块("p0","13800138000",{"页":0}),文本块("p1","Ordinary text stays",{"页":0})]
        detection=识别结果(出现=[出现记录("p0",0,11,"13800138000","mobile_phone",来源="规则")])
        output=[文本块("o0","[手机号1]",{"页":0}),文本块("o1","tays",{"页":0})]
        with self.assertRaises(ValueError):
            proc._核对实际写回(original,output,detection,全局映射表())


class _测试界面:
    def __init__(self):
        from types import SimpleNamespace,MethodType
        from 主程序 import app
        self.cancel_flag=False;self.is_processing=True;self._processor=None
        self.日志=[];self.状态=[];self.询问=[]
        控件=SimpleNamespace(configure=lambda **kw:None,set=lambda value:None)
        self.progressbar=self.progress_text=self.status_label=self.计时标签=控件
        self.run_desensitize=MethodType(app.脱敏工具GUI.run_desensitize,self)
    def after(self,delay,fn,*args):return fn(*args)
    def append_log(self,text):self.日志.append(text)
    def update_status_text(self,text):self.状态.append(text)
    def reset_ui(self):self.is_processing=False;self._processor=None
    def _启动计时(self):pass
    def _主线程询问(self,title,text):
        self.询问.append((title,text));return True


class GUI修复测试(unittest.TestCase):
    def 预检(self,输入,输出,ocr):
        from 主程序 import app
        from 主程序.ner引擎 import NER引擎
        from types import SimpleNamespace
        g=_测试界面();calls=[]
        def 不应请求(*args,**kwargs):
            calls.append(1)
            raise AssertionError("选择纯正则后不应再调用模型")
        boxes=SimpleNamespace(showwarning=lambda *args:None,showerror=lambda *args:None)
        with patch.object(app,"messagebox",boxes),patch.object(app,"探测lm_studio状态",return_value={"已启动":False,"模型已加载":False}),patch("主程序.脱敏处理器.默认ocr引擎",ocr),patch.object(NER引擎,"_请求文档模型",不应请求):
            app.脱敏工具GUI._预检并启动脱敏(g,str(输入),str(输出),False,False,False,False,False,"","","")
        self.assertEqual(calls,[])
        return g

    def test_文字PDF可选择纯正则并交付完整结果(self):
        import pymupdf
        from types import SimpleNamespace
        with 验证用临时目录() as d:
            p=Path(d)/"文字.pdf"
            with pymupdf.open() as doc:
                doc.new_page().insert_text((72,72),"Phone 13800138000");doc.save(p)
            out=Path(d)/"输出"
            g=self.预检(p,out,SimpleNamespace(是否可用=lambda:False))
            self.assertTrue(g.询问)
            files=list(out.glob("*.pdf"))
            self.assertEqual(len(files),1)
            with pymupdf.open(files[0]) as doc:
                self.assertIn("[手机号1]",doc[0].get_text())
                self.assertNotIn("13800138000",doc[0].get_text())
            self.assertIn("全部成功",g.状态[-1])

    def test_扫描混合批次不因缺模型扣住普通文件(self):
        import pymupdf
        from types import SimpleNamespace
        for available in (False,True):
            with self.subTest(OCR可用=available),验证用临时目录() as d:
                src=Path(d)/"输入";src.mkdir();out=Path(d)/"输出"
                (src/"普通.txt").write_text("电话13800138000",encoding="utf-8-sig")
                with pymupdf.open() as doc:
                    doc.new_page();doc.save(src/"扫描.pdf")
                ocr=SimpleNamespace(是否可用=lambda:available,pdf转文本=lambda p:"电话13900139000")
                g=self.预检(src,out,ocr)
                self.assertTrue((out/"普通_脱敏.txt").exists(), g.日志)
                self.assertEqual((out/"普通_脱敏.txt").read_text(encoding="utf-8-sig"),"电话[手机号1]")
                self.assertEqual((out/"扫描_脱敏.md").exists(),available)
                if available:self.assertNotIn("13900139000",(out/"扫描_脱敏.md").read_text(encoding="utf-8-sig"))
                self.assertTrue(g.询问)

    def test_已有文字层扫描PDF不因未装OCR阻断(self):
        import pymupdf
        from types import SimpleNamespace
        with 验证用临时目录() as d:
            text="Contact 13800138000. This searchable scan contains sufficient OCR words for local redaction."
            with pymupdf.open() as image:
                image.new_page().insert_text((30,72),text,fontsize=10)
                png=image[0].get_pixmap().tobytes("png")
            p=Path(d)/"扫描.pdf";out=Path(d)/"输出"
            with pymupdf.open() as doc:
                page=doc.new_page();page.insert_image(page.rect,stream=png)
                page.insert_text((30,72),text,fontsize=10,render_mode=3);doc.save(p)
            self.预检(p,out,SimpleNamespace(是否可用=lambda:False))
            self.assertEqual(len(list(out.glob("*.md"))),1)
            self.assertFalse(list(out.glob("*.pdf")))

    def test_取消批量还原显示取消与已处理比例(self):
        from 主程序 import app
        with 验证用临时目录() as d:
            files=[]
            for name in ("甲","乙"):
                p=Path(d)/(name+".txt");p.write_text("13800138000",encoding="utf-8-sig");files.append(str(p))
            proc=脱敏处理器(跳过模型=True,启用日期=False)
            result=proc.处理文件列表(files,str(Path(d)/"脱敏"))
            self.assertEqual(result.成功数,2)
            g=_测试界面();g._processor=proc;out=Path(d)/"还原"
            def 日志(text):
                g.日志.append(text)
                if text.startswith("还原文件已生成"):app.脱敏工具GUI.cancel_processing(g)
            g.append_log=日志
            app.脱敏工具GUI.run_restore(g,result.已生成文件,str(out))
            self.assertEqual(len(list(out.glob("*.txt"))),1)
            self.assertIn("取消",g.状态[-1])
            self.assertIn("1/2",g.状态[-1])
            self.assertNotIn("全部成功",g.状态[-1])


class 规则兼容边界测试(unittest.TestCase):
    def test_无冒号既有长账号继续识别(self):
        from 主程序.固定规则 import 固定规则识别
        text="收款银行账号1234567890123456789012。"
        found=固定规则识别(文本块("t",text,{}))
        self.assertIn(("1234567890123456789012","bank_account"),[(m.原词,m.类型) for m in found])

    def test_金额字段不污染同句后续数量(self):
        from 主程序.固定规则 import 固定规则识别
        text="金额：100元 数量：1234567件。"
        found=固定规则识别(文本块("t",text,{}),启用金额=True)
        self.assertFalse([m for m in found if m.类型 == "amount"])
