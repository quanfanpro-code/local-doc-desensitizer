from pathlib import Path
import importlib
import unittest
from 维护工具.运行脱敏验证 import 验证用临时目录
from 主程序.ner引擎 import 全局映射表
from 主程序.识别结果 import 出现记录, 识别结果


class 定位写回测试(unittest.TestCase):
    def test_PPT软换行与域文字的定位一致(self):
        from pptx import Presentation
        from pptx.util import Inches
        with 验证用临时目录() as 目录:
            p = Path(目录)/"原文.pptx"
            prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
            tf = slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(2)).text_frame
            tf.text = "成都星河\v科技有限公司"
            prs.save(p)
            out,back = self.处理(p,目录)
            self.assertEqual(Presentation(out).slides[0].shapes[0].text,"[公司1]\v")
            self.assertEqual(Presentation(back).slides[0].shapes[0].text,"成都星河\v科技有限公司")

    def test_PDF只改指定出现保存后可读并可还原(self):
        import pymupdf
        from 主程序.文档解析器 import 读取文档块
        from 主程序.格式保持器 import 保存定位脱敏, 还原并保存原格式
        with 验证用临时目录() as 目录:
            路径 = Path(目录)/"文字.pdf"
            with pymupdf.open() as doc:
                p = doc.new_page()
                p.insert_text((72,72),"BlueX and BlueX; phone 13800138000",fontsize=12)
                doc.save(路径)
            块 = 读取文档块(str(路径)); b = 块[0]
            结果 = 识别结果({"E1":{"canonical":None,"names":["BlueX"]}},
                [出现记录(b.编号,0,5,"BlueX","Ni","E1"),
                 出现记录(b.编号,23,34,"13800138000","mobile_phone")])
            映射 = 全局映射表()
            输出 = 保存定位脱敏(str(路径),块,结果,映射,目录)
            with pymupdf.open(输出) as doc:
                文本 = doc[0].get_text()
                self.assertEqual(文本.count("BlueX"),1)
                self.assertNotIn("13800138000",文本)
                self.assertIn("[机构1]",文本)
                self.assertIn("[手机号1]",文本)
                doc[0].get_pixmap(matrix=pymupdf.Matrix(1.5,1.5)).save(str(Path(目录)/"脱敏页.png"))
            mp = Path(目录)/"映射.json"; 映射.保存到文件(mp)
            重载 = 全局映射表(); 重载.从文件加载(mp)
            还原 = 还原并保存原格式(输出,重载,目录)
            with pymupdf.open(还原) as doc:
                文本 = doc[0].get_text()
                self.assertEqual(文本.count("BlueX"),2)
                self.assertIn("13800138000",文本)

    def test_Excel无缓存公式数字及引用往返(self):
        import hashlib
        from openpyxl import Workbook, load_workbook
        from 主程序.文档解析器 import 读取文档块
        from 主程序.格式保持器 import 保存定位脱敏, 还原并保存原格式
        with 验证用临时目录() as 目录:
            路径 = Path(目录)/"公式.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "星河"
            ws["A1"] = "成都星河科技有限公司"
            ws["B1"] = "='星河'!A1"
            ws["C1"] = '="成都星河科技有限公司"'
            ws["D1"] = "=SUM(1,2)"
            ws["E1"] = '="星"&"河"'
            ws["B2"] = 13800138000
            wb.save(路径)
            原哈希 = hashlib.sha256(路径.read_bytes()).digest()
            块 = 读取文档块(str(路径))
            cells = {b.位置["单元格"]:b for b in 块}
            self.assertEqual(cells["E1"].原文,"星河")
            self.assertEqual(cells["B1"].原文,"成都星河科技有限公司")
            结果 = 识别结果({"E1":{"canonical":"成都星河科技有限公司","names":["成都星河科技有限公司","星河"]}})
            for coord in ("A1","B1","C1","E1"):
                b = cells[coord]
                结果.出现.append(出现记录(b.编号,0,len(b.原文),b.原文,"Ni","E1"))
            b = cells["B2"]
            结果.出现.append(出现记录(b.编号,0,len(b.原文),b.原文,"mobile_phone"))
            映射 = 全局映射表()
            输出 = 保存定位脱敏(str(路径),块,结果,映射,目录)
            检查 = load_workbook(输出)
            self.assertEqual(检查.active["B1"].value,"='星河'!A1")
            self.assertEqual(检查.active["D1"].value,"=SUM(1,2)")
            self.assertEqual(检查.active["C1"].value,'="[公司1]"')
            self.assertEqual(检查.active["E1"].value,"[公司1]①")
            self.assertEqual(检查.active["B2"].value,"[手机号1]")
            检查.close()
            映射路径 = Path(目录)/"公式映射.json"; 映射.保存到文件(映射路径)
            重载 = 全局映射表(); 重载.从文件加载(映射路径)
            还原 = load_workbook(还原并保存原格式(输出,重载,目录))
            self.assertEqual(还原.active["E1"].value,'="星"&"河"')
            self.assertEqual(还原.active["B2"].value,13800138000)
            self.assertIsInstance(还原.active["B2"].value,int)
            self.assertEqual(原哈希,hashlib.sha256(路径.read_bytes()).digest())
            还原.close()

    def 处理(self, 路径, 目录):
        from 主程序 import 格式保持器
        from 主程序.文档解析器 import 读取文档块
        self.assertTrue(hasattr(格式保持器, "保存定位脱敏"), "尚未接入按位置写回")
        块 = 读取文档块(str(路径))
        命中 = []
        for b in 块:
            if "成都星河科技有限公司" in b.原文:
                s = b.原文.index("成都星河科技有限公司")
                命中.append(出现记录(b.编号,s,s+10,"成都星河科技有限公司","Ni","E1"))
        # 按字面长度修正测试记录，预期脱敏文本在下面独立给出。
        for m in 命中:
            m.结束 = m.开始 + len("成都星河科技有限公司")
        结果 = 识别结果({"E1":{"canonical":"成都星河科技有限公司","names":["成都星河科技有限公司","星河"]}},命中)
        映射 = 全局映射表()
        输出 = 格式保持器.保存定位脱敏(str(路径),块,结果,映射,str(目录))
        映射路径 = Path(目录)/"映射.json"
        映射.保存到文件(映射路径)
        重载 = 全局映射表(); 重载.从文件加载(映射路径)
        还原 = 格式保持器.还原并保存原格式(输出,重载,str(目录))
        return 输出,还原

    def test_Word跨样式只替换一次并准确还原文字及样式(self):
        from docx import Document
        with 验证用临时目录() as 目录:
            路径 = Path(目录)/"原文.docx"
            doc = Document()
            p = doc.add_paragraph()
            p.add_run("成都星河").bold = True
            p.add_run("科技有限公司").italic = True
            p.add_run("联合开展研发").underline = True
            doc.save(路径)
            输出,还原 = self.处理(路径,目录)
            p = Document(输出).paragraphs[0]
            self.assertEqual(p.text,"[公司1]联合开展研发")
            self.assertTrue(p.runs[-1].underline)
            p = Document(还原).paragraphs[0]
            self.assertEqual([r.text for r in p.runs],["成都星河","科技有限公司","联合开展研发"])
            self.assertTrue(p.runs[0].bold); self.assertTrue(p.runs[1].italic)

    def test_PPT跨片段与文本原文往返(self):
        from pptx import Presentation
        from pptx.util import Inches
        with 验证用临时目录() as 目录:
            路径 = Path(目录)/"原文.pptx"
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            p = slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(1)).text_frame.paragraphs[0]
            p.add_run().text = "成都星河"; p.add_run().text = "科技有限公司"
            p.runs[0].font.bold = True
            prs.save(路径)
            输出,还原 = self.处理(路径,目录)
            self.assertEqual(Presentation(输出).slides[0].shapes[0].text,"[公司1]")
            p = Presentation(还原).slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertEqual([r.text for r in p.runs],["成都星河","科技有限公司"])
            self.assertTrue(p.runs[0].font.bold)
        with 验证用临时目录() as 目录:
            路径 = Path(目录)/"原文.txt"
            路径.write_text("成都星河科技有限公司\n星河灿烂。\n",encoding="utf-8")
            输出,还原 = self.处理(路径,目录)
            self.assertEqual(Path(输出).read_text(encoding="utf-8-sig"),"[公司1]\n星河灿烂。\n")
            self.assertEqual(Path(还原).read_text(encoding="utf-8-sig"),"成都星河科技有限公司\n星河灿烂。\n")


class 文档块测试(unittest.TestCase):
    def test_短文字PDF不误当扫描件而混合扫描页不漏掉(self):
        import io
        import pymupdf
        from PIL import Image
        from 主程序.文档解析器 import 判断是否扫描版pdf
        with 验证用临时目录() as d:
            短 = Path(d)/"短文字.pdf"
            混合 = Path(d)/"混合页.pdf"
            with pymupdf.open() as doc:
                doc.new_page().insert_text((72,72),"BlueX")
                doc.save(短)
                page = doc.new_page()
                data = io.BytesIO(); Image.new("RGB",(100,100),"white").save(data,format="PNG")
                page.insert_image(page.rect,stream=data.getvalue())
                doc.save(混合)
            self.assertFalse(判断是否扫描版pdf(str(短)))
            self.assertTrue(判断是否扫描版pdf(str(混合)))

    def 读取(self, 路径):
        from 主程序 import 文档解析器
        self.assertTrue(callable(getattr(文档解析器, "读取文档块", None)), "缺少保留位置的文档提取")
        return 文档解析器.读取文档块(str(路径))

    def test_数字和混合金额单元格不丢失(self):
        from openpyxl import Workbook
        with 验证用临时目录() as 目录:
            路径 = Path(目录) / "供应商.xlsx"
            wb = Workbook()
            wb.active.append(["供应商", "手机号"])
            wb.active.append(["成都星河科技有限公司，金额100万元", 13800138000])
            wb.save(路径)
            块 = self.读取(路径)
        by_cell = {b.位置.get("单元格"): b for b in 块}
        self.assertEqual(by_cell["B2"].原文, "13800138000")
        self.assertEqual(by_cell["B2"].原值, 13800138000)
        self.assertIn("供应商", by_cell["A2"].上下文)
        self.assertIn("成都星河科技有限公司", by_cell["A2"].原文)

    def test_Word多文字节点及页眉均有独立位置(self):
        from docx import Document
        from docx.oxml import OxmlElement
        with 验证用临时目录() as 目录:
            路径 = Path(目录) / "样例.docx"
            doc = Document()
            run = doc.add_paragraph().add_run("成都星河")
            文字 = OxmlElement("w:t")
            文字.text = "科技有限公司"
            run._r.append(文字)
            doc.sections[0].header.paragraphs[0].text = "供应商简称星河"
            doc.save(路径)
            块 = self.读取(路径)
        self.assertIn("成都星河科技有限公司", [b.原文 for b in 块])
        self.assertIn("供应商简称星河", [b.原文 for b in 块])
        self.assertEqual(len({b.编号 for b in 块}), len(块))


class 完整入口测试(unittest.TestCase):
    def test_输出文件名不能覆盖同批另一份原件(self):
        from 主程序.脱敏处理器 import 脱敏处理器
        from unittest.mock import patch
        with 验证用临时目录() as d:
            a,b = Path(d)/"甲.txt",Path(d)/"甲_脱敏.txt"
            a.write_text("13800138000",encoding="utf-8")
            b.write_text("13900139000",encoding="utf-8")
            源字节 = [a.read_bytes(),b.read_bytes()]
            proc = 脱敏处理器()
            with patch.object(proc._ner,"识别文档",side_effect=lambda *a:识别结果()),patch.object(proc._ner,"复查输出",side_effect=lambda *a:识别结果()):
                result = proc.处理文件列表([str(a),str(b)],d)
            self.assertEqual(result.成功数,2)
            self.assertEqual([a.read_bytes(),b.read_bytes()],源字节)

    def test_文件之间映射隔离及同名输出不覆盖(self):
        from 主程序.脱敏处理器 import 脱敏处理器
        from unittest.mock import patch
        with 验证用临时目录() as d:
            files = []
            for name,text in (("甲","13800138000"),("乙","13900139000")):
                folder = Path(d)/name; folder.mkdir()
                file = folder/"原文.txt"; file.write_text(text,encoding="utf-8")
                files.append(str(file))
            proc = 脱敏处理器()
            with patch.object(proc._ner,"识别文档",side_effect=lambda *a:识别结果()),patch.object(proc._ner,"复查输出",side_effect=lambda *a:识别结果()):
                result = proc.处理文件列表(files,d)
            self.assertEqual(result.成功数,2)
            self.assertEqual(len(set(result.已生成文件)),2)
            for p in result.已生成文件:
                self.assertEqual(Path(p).read_text(encoding="utf-8-sig"),"[手机号1]")

    def test_扫描件走OCR到MD并准确还原(self):
        import pymupdf
        from unittest.mock import patch,Mock
        from 主程序.脱敏处理器 import 脱敏处理器
        with 验证用临时目录() as d:
            file = Path(d)/"扫描件.pdf"
            with pymupdf.open() as doc:
                doc.new_page(); doc.save(file)
            ocr = Mock(); ocr.是否可用.return_value = True
            ocr.pdf转文本.return_value = "电话13800138000"
            proc = 脱敏处理器(ocr引擎=ocr)
            with patch.object(proc._ner,"识别文档",return_value=识别结果()),patch.object(proc._ner,"复查输出",return_value=识别结果()):
                result = proc.处理文件列表([str(file)],d)
            self.assertEqual(result.成功数,1)
            self.assertTrue(result.已生成文件[0].endswith(".md"))
            restored = proc.还原文件列表(result.已生成文件,输出目录=d)
            self.assertEqual(Path(restored.成功文件[0]).read_text(encoding="utf-8-sig"),"电话13800138000")

    def test_局部识别失败仍交付文件及对应映射(self):
        from unittest.mock import patch
        from 主程序.脱敏处理器 import 脱敏处理器, 脱敏处理结果
        self.assertTrue(hasattr(脱敏处理结果(),"已生成文件"))
        with 验证用临时目录() as 目录:
            原件 = Path(目录)/"原文.txt"
            原件.write_text("供应商小麦，电话13800138000。",encoding="utf-8")
            处理器 = 脱敏处理器(启用日期=False)
            部分 = 识别结果(未完成=[{"块编号":"text:0","环节":"逐处核对","原因":"请求超时"}])
            with patch.object(处理器._ner,"识别文档",return_value=部分), patch.object(处理器._ner,"复查输出",return_value=识别结果()):
                结果 = 处理器.处理文件列表([str(原件)],目录)
            self.assertEqual(结果.成功数,0)
            self.assertEqual(结果.失败数,0)
            self.assertTrue(结果.未完成环节)
            self.assertEqual(len(结果.已生成文件),1)
            内容 = Path(结果.已生成文件[0]).read_text(encoding="utf-8-sig")
            self.assertIn("[手机号1]",内容)
            self.assertIn("小麦",内容)
            映射 = 全局映射表(); 映射.从文件加载(结果.映射表路径列表[0])
            self.assertEqual(映射.批量还原文本(内容),"供应商小麦，电话13800138000。")

    def test_实际输出复查发现新简称后补充定位写回(self):
        from unittest.mock import patch
        from 主程序.脱敏处理器 import 脱敏处理器, 脱敏处理结果
        self.assertTrue(hasattr(脱敏处理结果(),"已生成文件"))
        with 验证用临时目录() as 目录:
            原件 = Path(目录)/"原文.txt"
            原件.write_text("甲供应商向BlueX采购。",encoding="utf-8")
            处理器 = 脱敏处理器(启用日期=False)
            首次 = 识别结果({"E1":{"canonical":None,"names":["甲供应商"]}},
                [出现记录("text:0",0,4,"甲供应商","Ni","E1")],{"text:0"})
            def 复查(输出块,原块,已有机构,进度回调=None):
                b = 输出块[0]
                self.assertIn("[机构1]",b.原文)
                s = b.原文.index("BlueX")
                return 识别结果({"E2":{"canonical":None,"names":["BlueX"]}},
                    [出现记录(b.编号,s,s+5,"BlueX","Ni","E2")])
            with patch.object(处理器._ner,"识别文档",return_value=首次), patch.object(处理器._ner,"复查输出",side_effect=复查):
                结果 = 处理器.处理文件列表([str(原件)],目录)
            self.assertEqual(结果.成功数,1,结果.未完成环节)
            内容 = Path(结果.已生成文件[0]).read_text(encoding="utf-8-sig")
            self.assertEqual(内容,"[机构1]向[机构2]采购。")


class 验证入口测试(unittest.TestCase):
    def test_退出后保留材料并返回真实失败状态(self):
        try:
            入口 = importlib.import_module("维护工具.运行脱敏验证")
        except ImportError:
            self.fail("尚未提供可保留材料的验证入口")
        with 入口.验证用临时目录() as 目录:
            样例 = Path(目录) / "虚构样例.txt"
            样例.write_text("星河", encoding="utf-8-sig")
        self.assertEqual(样例.read_text(encoding="utf-8-sig"), "星河")
        用例 = unittest.FunctionTestCase(lambda: self.fail("用于核对退出状态的预期失败"))
        self.assertFalse(入口.运行测试(unittest.TestSuite([用例])).wasSuccessful())
